#!/usr/bin/env python3
"""Generate a professional PowerPoint summarizing the AMS-opt pipeline.

Creates a PPTX that explains:
- Structural graph (netlist -> SG)
- Functional/knowledge graph (LLM-generated KG, variant expansion)
- Combined graph fusion
- GNN feature vector definition
- GIN encoder + projection head architecture
- Losses: NT-Xent (+ weighted hard negatives) and consistency loss
- Example qualitative results (UMAP images if present)

Usage:
  python3 scripts/generate_presentation.py --out ams_opt_presentation.pptx

Optional:
  --umap-dir umap_results_full
"""

import argparse
import os
from datetime import date, datetime
from typing import Tuple


def _require_pptx():
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT  # type: ignore
        from pptx.util import Inches, Pt  # type: ignore
        from pptx.dml.color import RGBColor  # type: ignore

        return Presentation, PP_PARAGRAPH_ALIGNMENT, Inches, Pt, RGBColor
    except Exception as e:  # pragma: no cover
        raise RuntimeError(
            "Missing dependency 'python-pptx'. Install with: pip install python-pptx"
        ) from e


def _add_title_slide(prs, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def _add_bullets_slide(prs, title: str, bullets, notes=None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for i, b in enumerate(bullets):
        if i == 0:
            p = body.paragraphs[0]
        else:
            p = body.add_paragraph()
        p.text = b
        p.level = 0

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _add_two_column_slide(
    prs,
    title: str,
    left_title: str,
    left_bullets,
    right_title: str,
    right_bullets,
    notes=None,
) -> None:
    Presentation, PP_ALIGN, Inches, Pt, RGBColor = _require_pptx()

    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = title

    # Layout
    left_x = Inches(0.7)
    top_y = Inches(1.6)
    col_w = Inches(4.6)
    col_h = Inches(4.8)
    gap = Inches(0.4)
    right_x = left_x + col_w + gap

    # Left box
    left = slide.shapes.add_textbox(left_x, top_y, col_w, col_h)
    tf = left.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = left_title
    p0.font.bold = True
    p0.font.size = Pt(20)

    for b in left_bullets:
        p = tf.add_paragraph()
        p.text = b
        p.level = 1
        p.font.size = Pt(16)

    # Right box
    right = slide.shapes.add_textbox(right_x, top_y, col_w, col_h)
    tf2 = right.text_frame
    tf2.word_wrap = True
    q0 = tf2.paragraphs[0]
    q0.text = right_title
    q0.font.bold = True
    q0.font.size = Pt(20)

    for b in right_bullets:
        q = tf2.add_paragraph()
        q.text = b
        q.level = 1
        q.font.size = Pt(16)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _safe_add_picture(slide, image_path: str, left, top, width=None, height=None) -> bool:
    if not image_path or not os.path.exists(image_path):
        return False
    slide.shapes.add_picture(image_path, left, top, width=width, height=height)
    return True


def _add_results_slide(prs, umap_dir=None) -> None:
    Presentation, PP_ALIGN, Inches, Pt, RGBColor = _require_pptx()

    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "Qualitative Results (UMAP)"

    if not umap_dir:
        _add_bullets_slide(
            prs,
            "Qualitative Results (UMAP)",
            [
                "UMAP plots not provided to the generator.",
                "Pass --umap-dir umap_results_full to embed produced figures.",
            ],
        )
        return

    by_family = os.path.join(umap_dir, "umap_by_family.png")
    by_perf = os.path.join(umap_dir, "umap_by_performance.png")

    left = Inches(0.7)
    top = Inches(1.5)
    pic_w = Inches(6.0)

    ok1 = _safe_add_picture(slide, by_family, left, top, width=pic_w)
    ok2 = _safe_add_picture(slide, by_perf, Inches(7.0), top, width=pic_w)

    if not (ok1 or ok2):
        tf = slide.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(12.4), Inches(1.5)).text_frame
        tf.text = "UMAP images not found in the provided directory."
        tf.add_paragraph().text = f"Expected: {by_family} and/or {by_perf}"


def build_deck(out_path: str, umap_dir=None) -> None:
    Presentation, PP_ALIGN, Inches, Pt, RGBColor = _require_pptx()

    prs = Presentation()
    prs.core_properties.title = "AMS-opt: Circuit Graph Pipeline + Contrastive GNN"
    prs.core_properties.created = datetime.now()

    _add_title_slide(
        prs,
        title="AMS-opt: Analog Circuit Graph Pipeline + Contrastive GNN",
        subtitle=f"Structural + Functional graphs → Combined graph → GNN embeddings\nGenerated {date.today().isoformat()}",
    )

    _add_bullets_slide(
        prs,
        "Problem & Goal",
        [
            "Represent SPICE netlists as graphs suitable for learning.",
            "Fuse physical connectivity with functional (KG) knowledge.",
            "Learn circuit-level embeddings via contrastive learning.",
            "Enable retrieval/cluster analysis across circuit families.",
        ],
    )

    _add_bullets_slide(
        prs,
        "End-to-End Pipeline",
        [
            "1) Netlist → Structural Graph (`get_netlist_to_SG.py`)",
            "2) LLM KG → Functional Graph, then variant expansion (`transform_fun_graph.py`)",
            "3) Structural + Functional → Combined Graph (`combine_graphs.py`)",
            "4) Combined Graph → NPZ features/adjacency (`comb_graph_to_gnn.py`)",
            "5) Contrastive GNN training (`gnn_training/scripts/train.py`)",
        ],
    )

    _add_bullets_slide(
        prs,
        "Structural Graph (Stage 1)",
        [
            "Nodes: devices, terminals, nets.",
            "MOS is abstracted as device + (D,G,S) terminal nodes.",
            "Edges: device→terminal and terminal→net (MOS); device→net (generic).",
            "Output: `str_graph.json`.",
        ],
        notes=(
            "Implementation: get_netlist_to_SG.py\n"
            "MOS: dev:Mx, term:Mx:D/G/S, net:<name>.\n"
            "Non-MOS: dev:<name> directly connected to net nodes."
        ),
    )

    _add_bullets_slide(
        prs,
        "Functional / Knowledge Graph (Stage 2)",
        [
            "Nodes: performance metrics, substructures, parameters.",
            "Edges: relations like directly-proportional / trade-off / belongs-to / influences.",
            "Variant expansion creates relation-specific nodes.",
            "All relations are collapsed to `connects` edges between appropriate variants.",
            "Output: `fun_updated.json`.",
        ],
        notes=(
            "See scripts/transform_fun_graph.py\n"
            "Performance variants: -ambiguous, -trade-off, -directly-proportional\n"
            "Parameter variants: -directly-proportional, -inversely-proportional\n"
            "Links are rewritten to relation-conditioned variant endpoints; relation label becomes 'connects'."
        ),
    )

    _add_bullets_slide(
        prs,
        "Combined Graph Fusion (Stage 3)",
        [
            "Merge nodes and deduplicated links from SG and KG.",
            "Add device→parameter links for MOS devices: `dev:Mk`→`W_Mk`, `L_Mk`.",
            "Result is a single heterogeneous graph for learning.",
            "Output: `comb_graph.json`.",
        ],
        notes=(
            "See scripts/combine_graphs.py\n"
            "Merging strategy: functional nodes first, then overlay structural nodes (struct attrs win).\n"
            "Device-to-parameter links are created if missing."
        ),
    )

    _add_two_column_slide(
        prs,
        title="Node Feature Vector (Stage 4)",
        left_title="Type + Subcategory",
        left_bullets=[
            "Type one-hot (6): performance, sub-structure, parameter, net, device, terminal.",
            "Subcategory slots (4) depend on node type.",
            "Examples: performance {orig, ambiguous, trade-off, directly-prop}",
            "parameter {orig, directly-prop, inversely-prop, unused}",
        ],
        right_title="Meaning Block",
        right_bullets=[
            "Concatenated blocks: [perf meaning | substructure meaning].",
            "Perf uses base metric (variants map to base).",
            "Substructure IDs are normalized to canonical names.",
            "Produces `comb_graph_gnn.npz` + meta JSON.",
        ],
        notes=(
            "See scripts/comb_graph_to_gnn.py\n"
            "Feature layout: 6(type) + 4(subcat) + meaning_dim.\n"
            "Meaning dim is sized as max(perf_dim_max + sub_dim_max, 4).\n"
            "Performance 'Gain-trade-off' sets subcat slot but meaning uses 'Gain'."
        ),
    )

    _add_bullets_slide(
        prs,
        "GNN Architecture",
        [
            "Encoder: multi-layer GIN (sum aggregation + MLP).",
            "Readout: mean pooling over nodes → graph embedding.",
            "Projection head: MLP for contrastive loss.",
            "Outputs: embedding (encoder) and projection (used in loss).",
        ],
        notes=(
            "See gnn_training/models/gin_model.py\n"
            "GINConv: h' = MLP((1+ε)h + A·h)\n"
            "Readout: mean(h_nodes) then MLP to embedding_dim.\n"
            "ProjectionHead: 2-layer MLP to projection_dim."
        ),
    )

    _add_bullets_slide(
        prs,
        "Contrastive Objective: NT-Xent",
        [
            "Anchor z_i should match its positive z_j (same circuit, perturbed).",
            "Negatives: other positives in the batch; optional hard negatives.",
            "Similarity: cosine(z_i, z_j) / temperature.",
            "Loss: cross-entropy where the positive logit is the correct class.",
        ],
        notes=(
            "See gnn_training/losses/nt_xent_loss.py\n"
            "For each anchor k: logits = [sim(z_i[k], z_j[k]), sim(z_i[k], z_j[others]), hard_neg_sims].\n"
            "Label is always 0 (the first logit)."
        ),
    )

    _add_bullets_slide(
        prs,
        "Hard Negatives (Overlap-Based)",
        [
            "Mine negatives among circuits with non-zero substructure overlap.",
            "Weighted NT-Xent can upweight high-overlap negatives.",
            "This pushes apart circuits that look structurally similar.",
        ],
        notes=(
            "Hard negative miner lives under gnn_training/pairs/sampling.py\n"
            "Weighted loss: weights = 1 + hard_neg_weight * similarity_overlap."
        ),
    )

    _add_bullets_slide(
        prs,
        "Consistency Loss (Batch Similarity)",
        [
            "Compute pairwise circuit similarity matrix sim[i,j] ∈ [0,1] from substructure overlap.",
            "Encourage embedding distances to respect similarity structure.",
            "Two options: margin-based hinge or MSE on normalized distances.",
            "Combined loss: L = L_NT-Xent + w * L_consistency.",
        ],
        notes=(
            "See gnn_training/losses/consistency_loss.py\n"
            "Margin form uses: relu(margin + dist(i,j) - sim(i,j)*max_dist).\n"
            "Alternative uses MSE between normalized dist and (1-sim)."
        ),
    )

    _add_results_slide(prs, umap_dir)

    _add_bullets_slide(
        prs,
        "Artifacts & Repro",
        [
            "Graphs: `str_graph.json`, `fun_updated.json`, `comb_graph.json`.",
            "GNN data: `comb_graph_gnn.npz` + `comb_graph_gnn_meta.json`.",
            "Training: configs under `gnn_training/config/`, checkpoints under `checkpoints_*`.",
            "Viz: `scripts/visualize_embeddings_umap.py`, interactive plot HTML.",
        ],
    )

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    prs.save(out_path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="ams_opt_presentation.pptx", help="Output pptx path")
    ap.add_argument("--umap-dir", default=None, help="Directory with UMAP PNGs (e.g., umap_results_full)")
    args = ap.parse_args()

    build_deck(args.out, args.umap_dir)
    print(f"Wrote PowerPoint: {args.out}")


if __name__ == "__main__":
    main()
